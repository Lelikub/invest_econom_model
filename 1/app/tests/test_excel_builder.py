"""Tests for the formula-driven Excel Ground Truth workbook."""

from __future__ import annotations

import subprocess
import sys

from openpyxl import load_workbook

from src.comparator import compare_metrics
from src.excel_builder import (
    EXPECTED_RUSSIAN_SHEETS,
    build_workbook,
    read_excel_metrics,
    recalculate_with_excel,
)
from src.financial_model import FinancialModel
from src.monte_carlo import run_monte_carlo
from src.scenarios import run_boundary_tests, run_stress_test


def _build_real_workbook(real_project_data, path):
    base = FinancialModel(real_project_data).calculate()
    boundary = run_boundary_tests(real_project_data)
    monte_carlo = run_monte_carlo(real_project_data, iterations=1000, seed=20260916)
    stress = run_stress_test(real_project_data)
    cell_map = build_workbook(real_project_data, base, boundary, monte_carlo, stress, path)
    return base, cell_map


def test_workbook_contains_russian_sheets_and_linked_formulas(real_project_data, tmp_path):
    """Replacing formulas with Python values or English labels must fail."""

    path = tmp_path / "DCF_Ground_Truth.xlsx"
    _, cell_map = _build_real_workbook(real_project_data, path)
    workbook = load_workbook(path, data_only=False)

    assert workbook.sheetnames == EXPECTED_RUSSIAN_SHEETS
    dcf = workbook["DCF-модель"]
    metrics = workbook["Ключевые показатели"]
    assert dcf["B5"].value == "='Исходные данные'!$B$10"
    assert dcf["E5"].value == "=B5-C5-D5"
    assert dcf["I5"].value == "=F5+D5-G5-H5"
    assert metrics[cell_map.metrics["NPV"]].value.startswith("=NPV(")
    assert metrics[cell_map.metrics["IRR"]].value.startswith("=IRR(")
    assert workbook["Исходные данные"]["A1"].value == "Параметр"


def test_excel_com_recalculation_matches_python(real_project_data, tmp_path, capfd):
    """Uncalculated formulas or divergent Excel math must fail."""

    path = tmp_path / "DCF_Ground_Truth.xlsx"
    base, cell_map = _build_real_workbook(real_project_data, path)

    recalculate_with_excel(path)
    excel_metrics = read_excel_metrics(path, cell_map)
    comparison = compare_metrics(base.numeric_metrics(), excel_metrics, tolerance_percent=0.01)
    captured = capfd.readouterr()

    assert comparison.all_match
    assert comparison.max_relative_deviation_percent < 0.01
    assert "fatal exception" not in captured.err.lower()


def test_excel_com_shutdown_has_clean_stderr(real_project_data, tmp_path, assignment_root):
    """Releasing COM proxies after CoUninitialize must fail this test."""

    path = tmp_path / "DCF_Ground_Truth.xlsx"
    _build_real_workbook(real_project_data, path)
    code = (
        "import sys,site; "
        "sys.path.insert(0,r'app'); "
        "site.addsitedir(r'app\\.python-packages'); "
        "from src.excel_builder import recalculate_with_excel; "
        f"recalculate_with_excel(r'{path}')"
    )

    completed = subprocess.run(
        [sys.executable, "-c", code],
        cwd=assignment_root,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "fatal exception" not in completed.stderr.lower()
