"""End-to-end test for the one-command DCF/CAPEX pipeline."""

from __future__ import annotations

import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image

from src.pipeline import PipelineConfig, run_pipeline


def test_pipeline_creates_every_required_artifact(assignment_root, tmp_path):
    """Skipping a stage, gate, or final artifact must fail this test."""

    output_dir = tmp_path / "output"
    config = PipelineConfig(
        csv_path=assignment_root / "data" / "data.csv",
        output_dir=output_dir,
        log_path=tmp_path / "logs" / "execution.log",
    )

    result = run_pipeline(config)

    expected_names = {
        "DCF_Ground_Truth.xlsx",
        "Harness_Log.xlsx",
        "Сравнение_результатов.xlsx",
        "histogram_npv.png",
        "s_curve_npv.png",
        "Итоговый_отчет_ФИП.docx",
        "execution.log",
    }
    assert result.success
    assert result.comparison.all_match
    assert result.max_relative_deviation_percent < 0.01
    assert all(item.passed for item in result.boundary_tests)
    assert {path.name for path in result.artifacts} == expected_names
    for path in result.artifacts:
        assert path.exists() and path.stat().st_size > 0

    formula_book = load_workbook(output_dir / "DCF_Ground_Truth.xlsx", data_only=False, read_only=True)
    cached_book = load_workbook(output_dir / "DCF_Ground_Truth.xlsx", data_only=True, read_only=True)
    assert formula_book["Ключевые показатели"]["B4"].value.startswith("=NPV(")
    assert isinstance(cached_book["Ключевые показатели"]["B4"].value, (int, float))
    formula_book.close()
    cached_book.close()

    comparison_book = load_workbook(output_dir / "Сравнение_результатов.xlsx", data_only=True)
    assert {row[5].value for row in comparison_book.active.iter_rows(min_row=2)} == {"СОВПАДАЕТ"}
    comparison_book.close()

    report_path = output_dir / "Итоговый_отчет_ФИП.docx"
    assert zipfile.is_zipfile(report_path)
    assert Document(report_path).inline_shapes.__len__() == 2
    for chart_name in ("histogram_npv.png", "s_curve_npv.png"):
        with Image.open(output_dir / chart_name) as image:
            image.verify()

    log_text = config.log_path.read_text(encoding="utf-8")
    for stage in ("чтение CSV", "финансовая модель", "Excel", "Монте-Карло", "стресс-тест", "Word"):
        assert stage in log_text

