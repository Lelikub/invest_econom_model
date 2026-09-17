"""Tests for technical logs, Harness Log, and the generated Word report."""

from __future__ import annotations

from docx import Document
from openpyxl import load_workbook

from src.charts import save_monte_carlo_charts
from src.comparator import compare_metrics
from src.financial_model import FinancialModel
from src.harness_log import build_harness_entries, write_harness_log
from src.logger_config import configure_logging
from src.models import ReportContext
from src.monte_carlo import run_monte_carlo
from src.report_builder import build_word_report
from src.scenarios import run_boundary_tests, run_stress_test


def _report_context(data, directory):
    base = FinancialModel(data).calculate()
    boundary = run_boundary_tests(data)
    monte_carlo = run_monte_carlo(data, iterations=1000, seed=20260916)
    stress = run_stress_test(data)
    comparison = compare_metrics(base.numeric_metrics(), base.numeric_metrics())
    charts = save_monte_carlo_charts(monte_carlo, directory)
    entries = build_harness_entries(data, comparison, boundary)
    return ReportContext(data, base, boundary, monte_carlo, stress, comparison, charts, entries)


def test_runtime_logger_writes_utf8_stage_message(tmp_path):
    """A logger that omits the requested file or UTF-8 message must fail."""

    path = tmp_path / "execution.log"
    logger = configure_logging(path)
    logger.info("Проверка чтения CSV завершена")
    for handler in logger.handlers:
        handler.flush()

    content = path.read_text(encoding="utf-8")
    assert "INFO" in content
    assert "Проверка чтения CSV завершена" in content


def test_harness_contains_real_input_conflicts(real_project_data, tmp_path):
    """Omitting observed PDF/CSV discrepancies must fail."""

    context = _report_context(real_project_data, tmp_path)
    path = write_harness_log(context.harness_entries, tmp_path / "Harness_Log.xlsx")
    sheet = load_workbook(path, data_only=True).active
    values = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)

    assert sheet.cell(1, 1).value == "№"
    assert "WACC" in values and "10%" in values and "12%" in values
    assert "оборотного капитала" in values and "15%" in values
    assert "RPC_E_DISCONNECTED" in values


def test_word_report_contains_required_sections_images_and_calculated_values(real_project_data, tmp_path):
    """A hand-written or incomplete report must fail this artifact test."""

    context = _report_context(real_project_data, tmp_path)
    path = build_word_report(context, tmp_path / "Итоговый_отчет_ФИП.docx")
    document = Document(path)
    headings = {paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")}
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert {"Исходные данные", "Монте-Карло", "Стресс-тест", "Итоговый вывод"} <= headings
    assert len(document.inline_shapes) == 2
    assert len(document.tables) >= 6
    expected_npv = f"{context.base.npv:,.2f}".replace(",", " ").replace(".", ",")
    assert expected_npv in all_text
    assert "12,00%" in all_text
    assert path.stat().st_size > 50_000
