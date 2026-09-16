"""Build, recalculate, and read the formula-driven Excel Ground Truth."""

from __future__ import annotations

import math
import subprocess
import sys
from dataclasses import dataclass
from numbers import Real
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .models import (
    FinancialResult,
    MonteCarloResult,
    ProjectData,
    ScenarioTestResult,
    StressResult,
)


EXPECTED_RUSSIAN_SHEETS = [
    "Исходные данные",
    "Оборудование и CAPEX",
    "DCF-модель",
    "Ключевые показатели",
    "Сравнение Python Excel",
    "Граничные тесты",
    "Монте-Карло",
    "Стресс-тест",
    "Проверка исходных данных",
]


@dataclass(frozen=True)
class ExcelCellMap:
    """Stable locations of formula results read after Excel recalculation."""

    metrics: dict[str, str]
    year_metrics: dict[str, str]


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FORMULA_FILL = PatternFill("solid", fgColor="E2F0D9")
WARNING_FILL = PatternFill("solid", fgColor="FCE4D6")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)
MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'
PERCENT_FORMAT = "0.00%"


def build_workbook(
    data: ProjectData,
    base: FinancialResult,
    boundary_tests: tuple[ScenarioTestResult, ...],
    monte_carlo: MonteCarloResult,
    stress: StressResult,
    path: Path,
) -> ExcelCellMap:
    """Create a complete workbook whose core DCF results are Excel formulas."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {name: workbook.create_sheet(name) for name in EXPECTED_RUSSIAN_SHEETS}

    _write_inputs(sheets["Исходные данные"], data)
    equipment_rows = _write_capex(sheets["Оборудование и CAPEX"], data)
    year_metrics = _write_dcf(sheets["DCF-модель"], data, equipment_rows)
    metrics = _write_metrics(sheets["Ключевые показатели"], data, equipment_rows)
    _write_comparison_placeholder(sheets["Сравнение Python Excel"])
    _write_boundary_tests(sheets["Граничные тесты"], boundary_tests)
    _write_monte_carlo(sheets["Монте-Карло"], monte_carlo)
    _write_stress(sheets["Стресс-тест"], stress)
    _write_validations(sheets["Проверка исходных данных"], data)
    _add_defined_names(workbook, equipment_rows)
    _format_workbook(workbook)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(destination)
    return ExcelCellMap(metrics=metrics, year_metrics=year_metrics)


def recalculate_with_excel(path: Path) -> None:
    """Recalculate through an isolated COM worker and require clean completion."""

    target = Path(path).resolve()
    if not target.is_file():
        raise FileNotFoundError(f"Excel-файл не найден: {target}")
    worker = Path(__file__).with_name("excel_recalc_worker.py")
    creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    completed = subprocess.run(
        [sys.executable, str(worker), str(target)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
        creationflags=creation_flags,
    )
    stderr = completed.stderr.strip()
    if completed.returncode != 0 or stderr:
        details = stderr or completed.stdout.strip() or f"код возврата {completed.returncode}"
        raise RuntimeError(f"Microsoft Excel не смог чисто пересчитать книгу: {details}")



def read_excel_metrics(path: Path, cell_map: ExcelCellMap) -> dict[str, float | None]:
    """Read cached formula results after a real Excel recalculation."""

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        metrics_sheet = workbook["Ключевые показатели"]
        dcf_sheet = workbook["DCF-модель"]
        result: dict[str, float | None] = {}
        for name, coordinate in cell_map.metrics.items():
            result[name] = _coerce_cached_value(name, metrics_sheet[coordinate].value)
        for name, coordinate in cell_map.year_metrics.items():
            result[name] = _coerce_cached_value(name, dcf_sheet[coordinate].value)
        return result
    finally:
        workbook.close()


def write_comparison_sheet(path: Path, comparison) -> None:
    """Populate the workbook comparison sheet after cached results are read."""

    workbook = load_workbook(path)
    sheet = workbook["Сравнение Python Excel"]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row_index, item in enumerate(comparison.items, start=2):
        sheet.cell(row_index, 1, item.name)
        sheet.cell(row_index, 2, item.python_value)
        sheet.cell(row_index, 3, item.excel_value)
        sheet.cell(row_index, 4, item.absolute_difference)
        sheet.cell(row_index, 5, item.relative_difference_percent / 100 if item.relative_difference_percent is not None else None)
        sheet.cell(row_index, 6, item.status)
        for column in range(1, 7):
            sheet.cell(row_index, column).border = THIN_BORDER
        for column in (2, 3, 4):
            sheet.cell(row_index, column).number_format = MONEY_FORMAT
        sheet.cell(row_index, 5).number_format = PERCENT_FORMAT
    sheet.conditional_formatting.add(
        f"F2:F{max(2, sheet.max_row)}",
        CellIsRule(operator="equal", formula=['"НЕ СОВПАДАЕТ"'], fill=WARNING_FILL),
    )
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)


def write_comparison_workbook(comparison, path: Path) -> Path:
    """Save the Python-to-Excel comparison as a standalone workbook."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сравнение результатов"
    sheet.append(("Показатель", "Python", "Excel", "Абсолютное отклонение", "Относительное отклонение", "Статус"))
    for item in comparison.items:
        sheet.append(
            (
                item.name,
                item.python_value,
                item.excel_value,
                item.absolute_difference,
                None if item.relative_difference_percent is None else item.relative_difference_percent / 100,
                item.status,
            )
        )
        for column in (2, 3, 4):
            sheet.cell(sheet.max_row, column).number_format = MONEY_FORMAT
        sheet.cell(sheet.max_row, 5).number_format = PERCENT_FORMAT
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = (25, 18, 18, 23, 24, 18)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.conditional_formatting.add(
        f"F2:F{max(2, sheet.max_row)}",
        CellIsRule(operator="equal", formula=['"НЕ СОВПАДАЕТ"'], fill=WARNING_FILL),
    )
    workbook.save(destination)
    return destination



def _write_inputs(sheet, data: ProjectData) -> None:
    sheet.append(("Параметр", "Значение", "Единица / пояснение"))
    parameter_rows = [
        ("Коэффициент Ленга", data.lang_factor, "комбинированный solid-fluid процесс"),
        ("Доля оборотного капитала", data.working_capital_share, "доля TCI"),
        ("Ставка дисконтирования (WACC)", data.wacc, "доля"),
        ("Ставка налога на прибыль", data.tax_rate, "доля"),
        ("Срок проекта", data.project_lifetime_years, "лет"),
    ]
    for row in parameter_rows:
        sheet.append(row)
    sheet.append(())
    sheet.append(("Профили Revenue и OPEX из CSV",))
    sheet.append(("Год", "Выручка", "Операционные расходы"))
    for year in data.years:
        sheet.append((year, data.revenue[year], data.opex[year]))
    for row in range(2, 7):
        sheet.cell(row, 2).fill = INPUT_FILL
    for row in range(10, 10 + data.project_lifetime_years):
        sheet.cell(row, 2).fill = INPUT_FILL
        sheet.cell(row, 3).fill = INPUT_FILL
        sheet.cell(row, 2).number_format = MONEY_FORMAT
        sheet.cell(row, 3).number_format = MONEY_FORMAT
    for row in (3, 4, 5):
        sheet.cell(row, 2).number_format = PERCENT_FORMAT


def _write_capex(sheet, data: ProjectData) -> dict[str, int]:
    sheet.append(("Оборудование", "Поставочная стоимость"))
    for item in data.equipment:
        sheet.append((item.name, item.cost))
    equipment_total_row = len(data.equipment) + 2
    fci_row = equipment_total_row + 1
    wc_row = equipment_total_row + 2
    tci_row = equipment_total_row + 3
    sheet.cell(equipment_total_row, 1, "Сумма стоимости оборудования")
    sheet.cell(equipment_total_row, 2, f"=SUM(B2:B{equipment_total_row - 1})")
    sheet.cell(fci_row, 1, "Инвестиции в основной капитал (FCI)")
    sheet.cell(fci_row, 2, f"=B{equipment_total_row}*'Исходные данные'!$B$2")
    sheet.cell(wc_row, 1, "Доля оборотного капитала")
    sheet.cell(wc_row, 2, "='Исходные данные'!$B$3")
    sheet.cell(tci_row, 1, "Совокупные капитальные вложения (TCI)")
    sheet.cell(tci_row, 2, f"=B{fci_row}/(1-B{wc_row})")
    for row in range(2, tci_row + 1):
        sheet.cell(row, 2).number_format = PERCENT_FORMAT if row == wc_row else MONEY_FORMAT
        if row >= equipment_total_row:
            sheet.cell(row, 2).fill = FORMULA_FILL
    return {"equipment_total": equipment_total_row, "fci": fci_row, "wc": wc_row, "tci": tci_row}


def _write_dcf(sheet, data: ProjectData, equipment_rows: dict[str, int]) -> dict[str, str]:
    sheet["A1"] = "DCF-модель: год 0 содержит TCI, годы 1…N — только операционные потоки"
    headers = (
        "Год",
        "Выручка",
        "Операционные расходы",
        "Амортизация",
        "EBIT",
        "NOPAT",
        "Капитальные затраты",
        "Изменение оборотного капитала",
        "Свободный денежный поток",
        "Коэффициент дисконтирования",
        "Дисконтированный денежный поток",
        "Накопленный дисконтированный денежный поток",
        "Флаг окупаемости",
    )
    for column, value in enumerate(headers, start=1):
        sheet.cell(3, column, value)

    tci_ref = f"'Оборудование и CAPEX'!$B${equipment_rows['tci']}"
    fci_ref = f"'Оборудование и CAPEX'!$B${equipment_rows['fci']}"
    sheet.cell(4, 1, 0)
    for column in range(2, 9):
        sheet.cell(4, column, 0.0)
    sheet.cell(4, 9, f"=-{tci_ref}")
    sheet.cell(4, 10, 1.0)
    sheet.cell(4, 11, "=I4")
    sheet.cell(4, 12, "=K4")
    sheet.cell(4, 13, "=IF(L4>=0,1,0)")

    year_metrics: dict[str, str] = {}
    for offset, year in enumerate(data.years, start=1):
        row = 4 + offset
        input_row = 9 + year
        sheet.cell(row, 1, year)
        sheet.cell(row, 2, f"='Исходные данные'!$B${input_row}")
        sheet.cell(row, 3, f"='Исходные данные'!$C${input_row}")
        sheet.cell(row, 4, f"={fci_ref}/'Исходные данные'!$B$6")
        sheet.cell(row, 5, f"=B{row}-C{row}-D{row}")
        sheet.cell(row, 6, f"=E{row}*(1-'Исходные данные'!$B$5)")
        sheet.cell(row, 7, 0.0)
        sheet.cell(row, 8, 0.0)
        sheet.cell(row, 9, f"=F{row}+D{row}-G{row}-H{row}")
        sheet.cell(row, 10, f"=(1+'Исходные данные'!$B$4)^A{row}")
        sheet.cell(row, 11, f"=I{row}/J{row}")
        sheet.cell(row, 12, f"=L{row - 1}+K{row}")
        sheet.cell(row, 13, f"=IF(L{row}>=0,1,0)")
        year_metrics[f"Амортизация_{year}"] = f"D{row}"
        year_metrics[f"EBIT_{year}"] = f"E{row}"
        year_metrics[f"NOPAT_{year}"] = f"F{row}"
        year_metrics[f"FCF_{year}"] = f"I{row}"

    last_row = 4 + data.project_lifetime_years
    chart = LineChart()
    chart.title = "Денежные потоки проекта"
    chart.y_axis.title = "Условные денежные единицы"
    chart.x_axis.title = "Год"
    chart.add_data(Reference(sheet, min_col=9, min_row=3, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=last_row))
    chart.height = 8
    chart.width = 15
    sheet.add_chart(chart, "A17")
    return year_metrics


def _write_metrics(sheet, data: ProjectData, equipment_rows: dict[str, int]) -> dict[str, str]:
    sheet.append(("Показатель", "Значение", "Статус / метод"))
    last_dcf_row = 4 + data.project_lifetime_years
    metrics = {"FCI": "B2", "TCI": "B3", "NPV": "B4", "IRR": "B5", "PI": "B6", "DPBP": "B7"}
    rows = [
        (
            "Инвестиции в основной капитал (FCI)",
            f"='Оборудование и CAPEX'!B{equipment_rows['fci']}",
            "Метод Ленга",
        ),
        (
            "Совокупные капитальные вложения (TCI)",
            f"='Оборудование и CAPEX'!B{equipment_rows['tci']}",
            "FCI / (1 - WC)",
        ),
        (
            "Чистая приведённая стоимость (NPV)",
            f"=NPV('Исходные данные'!$B$4,'DCF-модель'!I5:I{last_dcf_row})+'DCF-модель'!I4",
            "Без терминальной стоимости",
        ),
        ("Внутренняя норма доходности (IRR)", f"=IRR('DCF-модель'!I4:I{last_dcf_row})", "Excel IRR"),
        (
            "Индекс рентабельности (PI)",
            f'=IF(B3=0,"НЕ ОПРЕДЕЛЕН",SUM(\'DCF-модель\'!K5:K{last_dcf_row})/B3)',
            "Стандартное определение PV будущих FCF / I0",
        ),
        (
            "Дисконтированный срок окупаемости (DPBP)",
            (
                f'=IF(MAX(\'DCF-модель\'!M5:M{last_dcf_row})=0,"НЕ ОКУПАЕТСЯ",'
                f'MATCH(1,\'DCF-модель\'!M5:M{last_dcf_row},0)-1+'
                f'(-INDEX(\'DCF-модель\'!L4:L{last_dcf_row - 1},MATCH(1,\'DCF-модель\'!M5:M{last_dcf_row},0)))'
                f'/INDEX(\'DCF-модель\'!K5:K{last_dcf_row},MATCH(1,\'DCF-модель\'!M5:M{last_dcf_row},0)))'
            ),
            "Линейная интерполяция внутри года",
        ),
    ]
    for row in rows:
        sheet.append(row)
    for row in range(2, 8):
        sheet.cell(row, 2).fill = FORMULA_FILL
        sheet.cell(row, 2).number_format = PERCENT_FORMAT if row == 5 else MONEY_FORMAT
    return metrics


def _write_comparison_placeholder(sheet) -> None:
    sheet.append(("Показатель", "Python", "Excel", "Абсолютное отклонение", "Относительное отклонение", "Статус"))
    sheet.append(("Сравнение заполняется после фактического пересчета Excel",))


def _write_boundary_tests(sheet, tests: tuple[ScenarioTestResult, ...]) -> None:
    sheet.append(("Код", "Тест", "Входные параметры", "Ожидаемое свойство", "Фактический результат", "Статус", "Комментарий"))
    for item in tests:
        sheet.append(
            (
                item.code,
                item.name,
                item.inputs,
                item.expected_property,
                item.actual_result,
                "PASS" if item.passed else "FAIL",
                item.comment,
            )
        )


def _write_monte_carlo(sheet, result: MonteCarloResult) -> None:
    sheet.append(("Параметр", "Значение"))
    summary = [
        ("Количество итераций", result.iterations),
        ("Random seed", result.seed),
        ("Распределение", "Логнормальное, среднее множителя 1,0"),
        ("Коэффициент вариации CAPEX", result.capex_cv),
        ("Коэффициент вариации Revenue", result.revenue_cv),
        ("Среднее NPV", result.mean),
        ("Медиана NPV", result.median),
        ("Стандартное отклонение NPV", result.standard_deviation),
        ("Минимум NPV", result.minimum),
        ("Максимум NPV", result.maximum),
        ("P(NPV < 0)", result.probability_negative),
    ]
    for item in summary:
        sheet.append(item)
    for level, value in result.quantiles.items():
        sheet.append((f"Квантиль {level:.0%}", value))
    data_start = sheet.max_row + 2
    sheet.cell(data_start, 1, "Итерация")
    sheet.cell(data_start, 2, "NPV")
    for index, value in enumerate(result.npv_values, start=1):
        sheet.cell(data_start + index, 1, index)
        sheet.cell(data_start + index, 2, float(value))
        sheet.cell(data_start + index, 2).number_format = MONEY_FORMAT
    for row in (5, 6, 12):
        sheet.cell(row, 2).number_format = PERCENT_FORMAT


def _write_stress(sheet, result: StressResult) -> None:
    sheet.append(("Показатель", "Базовый сценарий", "Стресс-сценарий", "Абсолютное изменение", "Относительное изменение"))
    for item in result.comparisons:
        sheet.append((item.name, item.base_value, item.stress_value, item.absolute_change, None if item.relative_change_percent is None else item.relative_change_percent / 100))
        for column in range(2, 5):
            sheet.cell(sheet.max_row, column).number_format = MONEY_FORMAT
        sheet.cell(sheet.max_row, 5).number_format = PERCENT_FORMAT
    sheet.cell(sheet.max_row + 2, 1, "Сценарий: CAPEX +100%, Revenue −30%")


def _write_validations(sheet, data: ProjectData) -> None:
    sheet.append(("Проверка", "Статус", "Подробности"))
    for record in data.validations:
        sheet.append((record.check, record.status, record.details))
    start = sheet.max_row + 2
    sheet.cell(start, 1, "Расхождения PDF и CSV")
    sheet.cell(start + 1, 1, "Параметр")
    sheet.cell(start + 1, 2, "PDF")
    sheet.cell(start + 1, 3, "CSV")
    sheet.cell(start + 1, 4, "Принято для расчета")
    for conflict in data.conflicts:
        sheet.append((conflict.field, conflict.pdf_value, conflict.csv_value, "CSV"))


def _add_defined_names(workbook: Workbook, equipment_rows: dict[str, int]) -> None:
    references = {
        "LANG_FACTOR": "'Исходные данные'!$B$2",
        "WORKING_CAPITAL_SHARE": "'Исходные данные'!$B$3",
        "WACC": "'Исходные данные'!$B$4",
        "TAX_RATE": "'Исходные данные'!$B$5",
        "PROJECT_LIFETIME": "'Исходные данные'!$B$6",
        "FCI": f"'Оборудование и CAPEX'!$B${equipment_rows['fci']}",
        "TCI": f"'Оборудование и CAPEX'!$B${equipment_rows['tci']}",
    }
    for name, reference in references.items():
        workbook.defined_names.add(DefinedName(name, attr_text=reference))


def _format_workbook(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "DCF-модель" else "A4"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            if cell.value is not None:
                cell.fill = HEADER_FILL
                cell.font = WHITE_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if sheet.title == "DCF-модель":
            for cell in sheet[3]:
                cell.fill = HEADER_FILL
                cell.font = WHITE_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.border = THIN_BORDER
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.fill = FORMULA_FILL
        for column in range(1, min(sheet.max_column, 13) + 1):
            max_length = max(
                (len(str(sheet.cell(row, column).value)) for row in range(1, min(sheet.max_row, 80) + 1) if sheet.cell(row, column).value is not None),
                default=8,
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(max(max_length + 2, 12), 34)
        sheet.auto_filter.ref = sheet.dimensions


def _coerce_cached_value(name: str, value) -> float | None:
    if isinstance(value, Real) and not isinstance(value, bool):
        number = float(value)
        if math.isfinite(number):
            return number
    if isinstance(value, str) and value in {"НЕ ОКУПАЕТСЯ", "НЕ ОПРЕДЕЛЕН", "#NUM!", "#N/A"}:
        return None
    if value is None:
        raise RuntimeError(f"Excel не сохранил вычисленное значение метрики {name}")
    raise RuntimeError(f"Некорректное cached value Excel для {name}: {value!r}")
