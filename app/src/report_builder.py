"""Generate the compact Russian-language investment report in Word."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

from .models import MetricOutcome, ReportContext


def build_word_report(context: ReportContext, path: Path) -> Path:
    """Build a report exclusively from finalized calculated results."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)
    _add_title_page(document)
    document.add_page_break()

    document.add_heading("Цель и постановка задачи", level=1)
    document.add_paragraph(
        "Цель работы — построить детерминированную DCF-модель R&D-центра синтеза "
        "твердотельных электролитов (TRL 3, комбинированный процесс), оценить CAPEX "
        "методом Ленга и проверить устойчивость результата. Расчеты выполнены Python, "
        "независимой формульной Excel Ground Truth моделью и зафиксированы в Harness Log."
    )

    document.add_heading("Исходные данные", level=1)
    document.add_paragraph(
        f"Горизонт модели — {context.data.project_lifetime_years} лет; WACC — {_percent(context.data.wacc)}; "
        f"налог на прибыль — {_percent(context.data.tax_rate)}; доля оборотного капитала — "
        f"{_percent(context.data.working_capital_share)}. Денежные значения представлены в условных "
        "единицах масштаба исходного CSV. Revenue и OPEX использованы как готовые агрегированные профили."
    )
    _add_table(
        document,
        ("Параметр", "PDF задания", "data.csv", "Принято"),
        [
            ("WACC", "10,00%", _percent(context.data.wacc), "data.csv"),
            ("Доля оборотного капитала", "12,00%", _percent(context.data.working_capital_share), "data.csv"),
            ("Коэффициент Ленга", "3,63 для комбинированного процесса", _number(context.data.lang_factor), "data.csv"),
        ],
    )
    _add_table(
        document,
        ("Оборудование", "Стоимость"),
        [(item.name, _money(item.cost)) for item in context.data.equipment]
        + [("Итого", _money(context.base.capital.equipment_cost))],
    )

    document.add_heading("Методика расчета", level=1)
    document.add_paragraph(
        "FCI = fL × ΣEi; TCI = FCI / (1 − WC). Амортизация линейная и рассчитывается от FCI; "
        "оборотный капитал не амортизируется. EBIT = Revenue − OPEX − D&A; NOPAT = EBIT × (1 − τ); "
        "FCF = NOPAT + D&A − CAPEX − ΔNWC. TCI учитывается один раз в году 0, а CAPEX и ΔNWC "
        "годов 1–N равны нулю из-за отсутствия отдельной динамики во входных данных. Терминальная "
        "стоимость не добавлялась. PI определен как PV будущих FCF / I0; DPBP интерполируется внутри года."
    )

    document.add_heading("Расчет CAPEX методом Ленга", level=1)
    document.add_paragraph(
        f"Сумма стоимости оборудования равна {_money(context.base.capital.equipment_cost)}. "
        f"Коэффициент Ленга {_number(context.data.lang_factor)} соответствует комбинированному "
        f"solid-fluid процессу. Получены FCI {_money(context.base.capital.fci)} и TCI "
        f"{_money(context.base.capital.tci)}."
    )

    document.add_heading("Результаты базовой DCF-модели", level=1)
    document.add_paragraph(f"NPV базового сценария: {_money(context.base.npv)}.")
    _add_table(
        document,
        ("Показатель", "Значение", "Статус"),
        [
            ("FCI", _money(context.base.capital.fci), "Рассчитано"),
            ("TCI", _money(context.base.capital.tci), "Рассчитано"),
            ("NPV", _money(context.base.npv), "Рассчитано"),
            ("IRR", _metric(context.base.irr, percent=True), context.base.irr.status),
            ("PI", _metric(context.base.pi), context.base.pi.status),
            ("DPBP", _metric(context.base.dpbp, suffix=" года"), context.base.dpbp.status),
        ],
    )
    document.add_page_break()

    document.add_heading("Проверка Excel Ground Truth и Python", level=1)
    document.add_paragraph(
        "Excel содержит собственные формулы со ссылками на входные ячейки, фактически пересчитан "
        "Microsoft Excel и затем прочитан обратно. Максимальное абсолютное отклонение равно "
        f"{_money(context.comparison.max_absolute_difference)}, максимальное относительное — "
        f"{_percent(context.comparison.max_relative_deviation_percent / 100)}. Требование < 0,01% "
        f"{'выполнено' if context.comparison.all_match else 'не выполнено'}."
    )
    _add_table(
        document,
        ("Показатель", "Python", "Excel", "Отклонение, %", "Статус"),
        [
            (
                item.name,
                _optional_number(item.python_value),
                _optional_number(item.excel_value),
                "—" if item.relative_difference_percent is None else _number(item.relative_difference_percent, 6),
                item.status,
            )
            for item in context.comparison.items[:6]
        ],
    )

    document.add_heading("Граничные тесты", level=1)
    _add_table(
        document,
        ("Тест", "Проверяемое свойство", "Фактический результат", "Статус"),
        [
            (item.name, item.expected_property, item.actual_result, "PASS" if item.passed else "FAIL")
            for item in context.boundary_tests
        ],
    )
    document.add_page_break()

    document.add_heading("Монте-Карло", level=1)
    document.add_paragraph(
        f"Выполнено {context.monte_carlo.iterations} итераций с seed {context.monte_carlo.seed}. "
        "Применены независимые логнормальные множители со средним 1,0: коэффициент вариации CAPEX "
        f"{_percent(context.monte_carlo.capex_cv)}, Revenue {_percent(context.monte_carlo.revenue_cv)}. "
        "Такой выбор исключает отрицательный CAPEX. S-Curve показывает P(NPV ≤ x)."
    )
    _add_table(
        document,
        ("Статистика", "Значение"),
        [
            ("Среднее NPV", _money(context.monte_carlo.mean)),
            ("Медиана NPV", _money(context.monte_carlo.median)),
            ("Стандартное отклонение", _money(context.monte_carlo.standard_deviation)),
            ("Квантиль 5%", _money(context.monte_carlo.quantiles[0.05])),
            ("Квантиль 95%", _money(context.monte_carlo.quantiles[0.95])),
            ("P(NPV < 0)", _percent(context.monte_carlo.probability_negative)),
        ],
    )
    image_table = document.add_table(rows=1, cols=2)
    for cell, chart, caption in zip(
        image_table.rows[0].cells,
        context.charts,
        ("Рисунок 1 — Гистограмма NPV", "Рисунок 2 — S-Curve P(NPV ≤ x)"),
    ):
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(chart), width=Inches(3.0))
        caption_paragraph = cell.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("Стресс-тест", level=1)
    document.add_paragraph("Сценарий: CAPEX +100%, Revenue −30%. Остальные предпосылки сохранены.")
    _add_table(
        document,
        ("Показатель", "Базовый", "Стресс", "Абсолютное изменение", "Относительное изменение"),
        [
            (
                item.name,
                _optional_number(item.base_value),
                _optional_number(item.stress_value),
                _optional_number(item.absolute_change),
                "—" if item.relative_change_percent is None else f"{_number(item.relative_change_percent)}%",
            )
            for item in context.stress.comparisons[:6]
        ],
    )

    document.add_heading("Harness Log и выявленные ошибки", level=1)
    _add_table(
        document,
        ("Этап", "Проблема", "Исправление", "Повторная проверка"),
        [
            (entry.stage, entry.detected_issue, entry.correction, entry.recheck_result)
            for entry in context.harness_entries
        ],
    )

    document.add_heading("Итоговый вывод", level=1)
    document.add_paragraph(_conclusion(context))
    document.core_properties.title = "Итоговый отчет ФИП: DCF и факторная оценка CAPEX"
    document.core_properties.subject = "R&D-центр синтеза твердотельных электролитов"
    document.save(destination)
    return destination


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _add_title_page(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.space_after = Pt(24)
    run = paragraph.add_run("DCF-моделирование и факторная оценка CAPEX в High-Tech")
    run.bold = True
    run.font.size = Pt(18)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Итоговый инвестиционный отчет\n").bold = True
    subtitle.add_run("R&D-центр синтеза твердотельных электролитов\nTRL 3, комбинированный процесс")
    document.add_paragraph()
    tools = document.add_paragraph("Инструменты: Python, Excel Ground Truth, Harness Log")
    tools.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_table(document: Document, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def _money(value: float) -> str:
    return _number(value, 2)


def _number(value: float, digits: int = 2) -> str:
    return f"{value:,.{digits}f}".replace(",", " ").replace(".", ",")


def _percent(value: float) -> str:
    return f"{_number(value * 100, 2)}%"


def _optional_number(value: float | None) -> str:
    return "Не определено" if value is None else _number(value, 6)


def _metric(outcome: MetricOutcome, *, percent: bool = False, suffix: str = "") -> str:
    if outcome.value is None:
        return "Не определено"
    if percent:
        return _percent(outcome.value)
    return f"{_number(outcome.value)}{suffix}"


def _conclusion(context: ReportContext) -> str:
    base_assessment = (
        "Базовый сценарий создает стоимость"
        if context.base.npv > 0
        else "Базовый сценарий не создает стоимость при принятой ставке дисконтирования"
    )
    stress_assessment = (
        "стресс-сценарий сохраняет положительный NPV"
        if context.stress.stress.npv > 0
        else "стресс-сценарий приводит к отрицательному NPV"
    )
    return (
        f"{base_assessment}: NPV {_money(context.base.npv)}. По Монте-Карло вероятность отрицательного "
        f"NPV составляет {_percent(context.monte_carlo.probability_negative)}; {stress_assessment} "
        f"({_money(context.stress.stress.npv)}). Excel и Python {'согласованы' if context.comparison.all_match else 'не согласованы'} "
        f"при максимальном относительном отклонении {_number(context.comparison.max_relative_deviation_percent, 6)}%. "
        "Вывод ограничен входным горизонтом, отсутствием терминальной стоимости и агрегированными Revenue/OPEX."
    )

