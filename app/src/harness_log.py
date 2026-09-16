"""Create the auditable Harness Log from observed issues and decisions."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import ComparisonResult, HarnessEntry, ProjectData, ScenarioTestResult


HEADERS = (
    "№",
    "Этап",
    "Запрос к ИИ / задача",
    "Сгенерированный или изменённый фрагмент кода",
    "Обнаруженная ошибка / галлюцинация",
    "Причина",
    "Как исправлено",
    "Результат повторной проверки",
)


def build_harness_entries(
    data: ProjectData,
    comparison: ComparisonResult,
    boundary_tests: tuple[ScenarioTestResult, ...],
) -> tuple[HarnessEntry, ...]:
    """Build entries for issues actually observed during this implementation."""

    conflicts = {item.field: item for item in data.conflicts}
    entries: list[HarnessEntry] = []
    if "wacc" in conflicts:
        entries.append(
            HarnessEntry(
                1,
                "Аудит исходных данных",
                "Сопоставить WACC задания и data.csv",
                'PDF_CONTROL_VALUES["wacc"] = 0.10',
                "WACC: PDF 10%, data.csv 12%",
                "Исходные материалы содержат разные сценарные значения",
                "Оба значения раскрыты; Python и Excel используют машинный вход CSV 12%",
                "Расхождение отражено в Excel, логе и отчете; расчеты используют 12%",
            )
        )
    if "working_capital_share" in conflicts:
        entries.append(
            HarnessEntry(
                len(entries) + 1,
                "Аудит исходных данных",
                "Сопоставить долю оборотного капитала задания и data.csv",
                'PDF_CONTROL_VALUES["working_capital_share"] = 0.12',
                "Доля оборотного капитала: PDF 12%, data.csv 15%",
                "Исходные материалы содержат разные сценарные значения",
                "Оба значения раскрыты; Python и Excel используют машинный вход CSV 15%",
                "TCI в обеих моделях рассчитан как FCI / (1 − 15%)",
            )
        )
    entries.extend(
        [
            HarnessEntry(
                len(entries) + 1,
                "Методология PI",
                "Проверить наличие явной формулы PI в лекции",
                "PI = PV будущих FCF / I0",
                "В лекции PI упомянут как метрика, но отдельная формула не приведена",
                "Исходный материал не определяет расчет PI",
                "Применено стандартное финансовое определение с явной пометкой допущения",
                "Python и Excel используют одну раскрытую формулу; сверка пройдена",
            ),
            HarnessEntry(
                len(entries) + 2,
                "Архитектура денежных потоков",
                "Исключить двойной учет CAPEX и оборотного капитала",
                "cash_flows = (-TCI, *annual_fcf); annual CAPEX = 0; ΔNWC = 0",
                "Риск повторного вычитания TCI в операционных годах",
                "TCI уже полностью признан оттоком года 0, иной динамики входные данные не содержат",
                "CAPEX и ΔNWC годов 1–N установлены в 0; амортизируется только FCI",
                "Unit-тест подтверждает единственный инвестиционный отток года 0",
            ),
            HarnessEntry(
                len(entries) + 3,
                "Пересчет Excel",
                "Выполнить фактический COM-пересчет формул",
                "Excel COM вынесен в excel_recalc_worker.py",
                "Первый интеграционный запуск вывел RPC_E_DISCONNECTED при CoUninitialize",
                "COM-proxy освобождались внутри долгоживущего pytest-процесса",
                "Пересчет изолирован в отдельном короткоживущем worker-процессе",
                "Полный тест проходит с чистым stderr; cached values прочитаны",
            ),
        ]
    )
    entries.append(
        HarnessEntry(
            len(entries) + 1,
            "Сверка Python и Excel",
            "Сравнить расчетные результаты при допуске < 0,01%",
            "compare_metrics(..., tolerance_percent=0.01)",
            "Расхождения выше допуска не обнаружены" if comparison.all_match else "Обнаружены расхождения выше допуска",
            "Независимые реализации используют одинаковые раскрытые исходные данные",
            "Сравнены FCI, TCI, метрики и годовые DCF-строки",
            (
                f"Максимальное относительное отклонение {comparison.max_relative_deviation_percent:.8f}%"
                if comparison.all_match
                else "Pipeline заблокирован до устранения расхождений"
            ),
        )
    )
    failed_cases = [item.code for item in boundary_tests if not item.passed]
    if failed_cases:
        entries.append(
            HarnessEntry(
                len(entries) + 1,
                "Граничные тесты",
                "Запустить обязательные математические сценарии",
                "run_boundary_tests(data)",
                f"Не прошли сценарии: {', '.join(failed_cases)}",
                "Нарушено ожидаемое математическое свойство",
                "Требуется исправление модели до формирования отчета",
                "FAIL",
            )
        )
    return tuple(entries)


def write_harness_log(entries: tuple[HarnessEntry, ...], path: Path) -> Path:
    """Write the structured Harness Log as a formatted XLSX workbook."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Harness Log"
    sheet.append(HEADERS)
    for entry in entries:
        sheet.append(
            (
                entry.number,
                entry.stage,
                entry.task,
                entry.code_fragment,
                entry.detected_issue,
                entry.cause,
                entry.correction,
                entry.recheck_result,
            )
        )
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="B7B7B7")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = (6, 24, 32, 40, 40, 35, 42, 38)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    workbook.save(destination)
    return destination

